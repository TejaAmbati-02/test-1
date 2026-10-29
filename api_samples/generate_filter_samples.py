"""Generate filter request/response samples (XML + Bruno collection) from sample_products.json.

The responses are what POST /v1/search returns when the index holds only the
products in sample_products.json (taken from TejaAmbati-02/Test_Data
out_transformed). Filter validation, matching, facets, sorting, pagination and
the product shape follow routers/search.py and services/search.py.

Usage:
    python api_samples/generate_filter_samples.py
"""
from __future__ import annotations

import json
import uuid
from pathlib import Path
from xml.sax.saxutils import escape

HERE = Path(__file__).parent
PRODUCTS = json.loads((HERE / "sample_products.json").read_text(encoding="utf-8"))

# Same as services/search.py
FILTER_KEY_MAP = {
    "category": "category_ids", "sub_category": "sub_category_id", "seller": "seller_details.seller_id",
    "category_name": "category_ids", "sub_category_name": "sub_category_id", "sellername": "seller_details.seller_id",
    "lead_time": "seller_details.lead_time", "category_ids": "category_ids", "sub_category_id": "sub_category_id",
    "product_type_id": "product_type_id", "product_type": "product_type.keyword", "material_type": "material_type",
    "availability": "availability", "price": "price", "moq": "moq", "currency": "currency",
    "sold_by": "sold_by.keyword", "supplier_count": "supplier_count", "contract_indicator": "contract_indicator",
    "temporary_unavailable": "temporary_unavailable", "sku_id": "sku_id", "product_key": "product_key",
    "asmo_material_number": "asmo_material_number", "aramco_material_number": "aramco_material_number",
}
FILTER_ALT_FIELDS = {"product_type": ("product_type_id", "product_type.keyword")}
FACETS = {
    "category": ("category_ids", "category_name"),
    "sub_category": ("sub_category_id", "sub_category_name"),
    "product_type": ("product_type_id", "product_type"),
    "seller": ("seller_details.seller_id", "seller_details.seller_name"),
}
# Same as routers/search.py
NUMERIC_FILTER_KEYS = {"price", "moq", "supplier_count", "lead_time"}
BOOLEAN_FILTER_KEYS = {"contract_indicator", "temporary_unavailable"}
RANGE_OPERATORS = {"gt", "gte", "lt", "lte"}


class ApiError(Exception):
    def __init__(self, status: int, code: str, message: str) -> None:
        self.status, self.body = status, {"code": code, "message": message}


# ---- validation (routers/search.py: _parse_filters / _validate_filter_value / _build_os_filters)

def _valid_scalar(key, value):
    if key in NUMERIC_FILTER_KEYS:
        return isinstance(value, (int, float)) and not isinstance(value, bool)
    if key in BOOLEAN_FILTER_KEYS:
        return isinstance(value, bool)
    return isinstance(value, str) and value != ""


def parse_filters(raw):
    if not raw or not raw.strip():
        return {}
    try:
        parsed = json.loads(raw)
    except json.JSONDecodeError as exc:
        raise ApiError(400, "INVALID_FILTER", f"filters must be a valid JSON object: {exc.msg}")
    if not isinstance(parsed, dict):
        raise ApiError(400, "INVALID_FILTER", "filters must be a JSON object")
    for k, v in parsed.items():
        if k not in FILTER_KEY_MAP:
            raise ApiError(400, "INVALID_FILTER_KEY", f"The filter key '{k}' is not supported.")
        if isinstance(v, dict):
            ok = k in NUMERIC_FILTER_KEYS and v and set(v) <= RANGE_OPERATORS and all(_valid_scalar(k, x) for x in v.values())
        elif isinstance(v, list):
            ok = bool(v) and all(_valid_scalar(k, x) for x in v)
        else:
            ok = _valid_scalar(k, v)
        if not ok:
            raise ApiError(400, "INVALID_FILTER_VALUE", f"Invalid value for filter '{k}': {json.dumps(v)}")
    return parsed


# ---- matching (term / terms / range, nested seller_details, product_type id-or-name)

def _values(doc, os_field):
    field = os_field.removesuffix(".keyword")
    if field.startswith("seller_details."):
        return [s.get(field.split(".", 1)[1]) for s in doc.get("seller_details") or []]
    v = doc.get(field)
    return v if isinstance(v, list) else [v]


def _match_value(value, cond):
    if value is None:
        return False
    if isinstance(cond, dict):
        ops = {"gt": value.__gt__, "gte": value.__ge__, "lt": value.__lt__, "lte": value.__le__}
        return all(ops[op](x) for op, x in cond.items())
    return any(_eq(value, c) for c in (cond if isinstance(cond, list) else [cond]))


def _eq(a, b):
    # True == 1 in Python, but a boolean term never matches a number (and vice versa).
    return a == b and isinstance(a, bool) == isinstance(b, bool)


def matches(doc, filters):
    for k, cond in filters.items():
        fields = FILTER_ALT_FIELDS.get(k, (FILTER_KEY_MAP[k],))
        if not any(_match_value(v, cond) for f in fields for v in _values(doc, f)):
            return False
    return True


# ---- response shaping (services/search.py: _to_product / _parse_facets)

def to_product(doc):
    sellers = doc.get("seller_details") or []
    return {
        "sku_id": doc.get("sku_id"),
        "sku_name": doc.get("sku_name") or doc.get("product_name"),
        "primary_image_url": doc.get("primary_image_url"),
        "asmo_material_number": doc.get("asmo_material_number"),
        "aramco_material_number": doc.get("aramco_material_number"),
        "seller_details": [{
            "seller_id": s.get("seller_id"),
            "name": s.get("seller_name"),
            "lead_time": s.get("lead_time"),
            "unit_type": s.get("unit_type"),
            "price": s.get("price"),
            "moq": s["moq"] if s.get("moq") is not None else "",
            "currency": s.get("currency"),
            "temporary_unavailable": s["temporary_unavailable"] if s.get("temporary_unavailable") is not None
            else (doc.get("temporary_unavailable") if i == 0 else None),
        } for i, s in enumerate(sellers)],
    }


def facets(docs):
    out = {}
    for key, (id_field, label_field) in FACETS.items():
        buckets: dict[str, dict] = {}
        for d in docs:
            if id_field.startswith("seller_details."):
                pairs = {(s.get("seller_id"), s.get("seller_name")) for s in d.get("seller_details") or []}
            else:
                ids = d.get(id_field)
                pairs = {(i, d.get(label_field)) for i in (ids if isinstance(ids, list) else [ids])}
            for i, label in pairs:  # count products, not seller offers
                if i is None:
                    continue
                b = buckets.setdefault(i, {"id": i, "value": label if label is not None else i, "count": 0})
                b["count"] += 1
        out[key] = sorted(buckets.values(), key=lambda b: (-b["count"], b["id"]))
    return out


def search(case_id, form):
    """Run one POST /v1/search filter-browse request against the sample products."""
    if form.get("channel") not in ("web", "mobile"):
        raise ApiError(400, "UNSUPPORTED_CHANNEL", f"Unsupported channel '{form.get('channel')}'. Allowed: mobile, web")
    sort_by, sort_order = form.get("sort_by"), form.get("sort_order", "desc")
    if sort_by is not None and sort_by not in ("relevance", "price"):
        raise ApiError(400, "INVALID_SORT_BY", f"Unsupported sort_by '{sort_by}'. Allowed: price, relevance")
    page, page_size = int(form.get("page", 1)), int(form.get("page_size", 60))
    filters = parse_filters(form.get("filters"))
    if not filters:
        raise ApiError(400, "EMPTY_REQUEST", "At least one of search_query, image, filters or request_id must be provided")
    hits = [d for d in PRODUCTS if matches(d, filters)]
    if sort_by == "price":  # documents without a price sort last in both directions
        priced = sorted((d for d in hits if d.get("price") is not None), key=lambda d: d["price"], reverse=sort_order == "desc")
        hits = priced + [d for d in hits if d.get("price") is None]
    start = (page - 1) * page_size
    return {
        "request_id": str(uuid.uuid5(uuid.NAMESPACE_URL, f"asmo-filter-sample/{case_id}")),
        "total_results": len(hits),
        "page": page,
        "page_size": page_size,
        "products": [to_product(d) for d in hits[start:start + page_size]],
        "facets": facets(hits),
    }


# ---- cases

P = {d["sku_id"]: d for d in PRODUCTS}
MCC, CORR, NOTE = P["038500-001-electrical"], P["001000-004-chemicals-1"], P["a-880000200002_1"]
HPE = P["a-880000629768"]
CAT_ELEC, CAT_IT = MCC["category_ids"][0], HPE["category_ids"][0]
SELLER_WW, SELLER_GCC = "3a7b23e0-0047-4d63-88a9-cd05d1006a39", "7f67afea-9853-43b5-bc5b-69904e4e2a6f"

# (id, folder, name, filters, extra form fields)
CASES: list[tuple[str, str, str, object, dict]] = [
    # Facet filters (send back the facet id) and their accepted aliases
    ("F01", "1. Category filters", "category - single id", {"category": [CAT_IT]}, {}),
    ("F02", "1. Category filters", "category - list of ids (match any)", {"category": [CAT_IT, CAT_ELEC]}, {}),
    ("F03", "1. Category filters", "category - scalar id", {"category": CAT_ELEC}, {}),
    ("F04", "1. Category filters", "category_name (alias of category)", {"category_name": [CAT_IT]}, {}),
    ("F05", "1. Category filters", "category_ids", {"category_ids": [CAT_ELEC]}, {}),
    ("F06", "2. Sub-category filters", "sub_category - id", {"sub_category": [HPE["sub_category_id"]]}, {}),
    ("F07", "2. Sub-category filters", "sub_category_name (alias of sub_category)", {"sub_category_name": [CORR["sub_category_id"]]}, {}),
    ("F08", "2. Sub-category filters", "sub_category_id - scalar", {"sub_category_id": MCC["sub_category_id"]}, {}),
    ("F09", "3. Product-type filters", "product_type - facet id", {"product_type": [HPE["product_type_id"]]}, {}),
    ("F10", "3. Product-type filters", "product_type - name", {"product_type": [MCC["product_type"]]}, {}),
    ("F11", "3. Product-type filters", "product_type - id and name mixed (match any)", {"product_type": [HPE["product_type_id"], CORR["product_type"]]}, {}),
    ("F12", "3. Product-type filters", "product_type_id", {"product_type_id": NOTE["product_type_id"]}, {}),
    ("F13", "4. Seller filters", "seller - GCC Control id", {"seller": [SELLER_GCC]}, {}),
    ("F14", "4. Seller filters", "seller - both seller ids (match any)", {"seller": [SELLER_GCC, SELLER_WW]}, {}),
    ("F15", "4. Seller filters", "sellername (alias of seller)", {"sellername": [SELLER_GCC]}, {}),
    ("F16", "4. Seller filters", "sold_by - name", {"sold_by": "GCC Control"}, {}),
    ("F17", "4. Seller filters", "sold_by - list", {"sold_by": ["GCC Control", "Wholesale Warehouse"]}, {}),
    # Keyword filters
    ("F18", "5. Keyword filters", "material_type - 9COM", {"material_type": "9COM"}, {}),
    ("F19", "5. Keyword filters", "material_type - list 9CAT or 9COM", {"material_type": ["9CAT", "9COM"]}, {}),
    ("F20", "5. Keyword filters", "availability - IN_STOCK", {"availability": "IN_STOCK"}, {}),
    ("F21", "5. Keyword filters", "availability - OUT_OF_STOCK (no match)", {"availability": ["OUT_OF_STOCK"]}, {}),
    ("F22", "5. Keyword filters", "currency - USD", {"currency": "USD"}, {}),
    ("F23", "5. Keyword filters", "currency - list SAR or USD", {"currency": ["SAR", "USD"]}, {}),
    ("F24", "5. Keyword filters", "sku_id - single", {"sku_id": HPE["sku_id"]}, {}),
    ("F25", "5. Keyword filters", "sku_id - list", {"sku_id": [HPE["sku_id"], MCC["sku_id"]]}, {}),
    ("F26", "5. Keyword filters", "product_key", {"product_key": CORR["product_key"]}, {}),
    ("F27", "5. Keyword filters", "asmo_material_number", {"asmo_material_number": MCC["asmo_material_number"]}, {}),
    ("F28", "5. Keyword filters", "aramco_material_number", {"aramco_material_number": HPE["aramco_material_number"]}, {}),
    # Numeric filters: exact, list, every range operator, combined ranges
    ("F29", "6. Price filters", "price - exact 80", {"price": 80}, {}),
    ("F30", "6. Price filters", "price - list [53.5, 20.5]", {"price": [53.5, 20.5]}, {}),
    ("F31", "6. Price filters", "price - gt 80", {"price": {"gt": 80}}, {}),
    ("F32", "6. Price filters", "price - gte 80", {"price": {"gte": 80}}, {}),
    ("F33", "6. Price filters", "price - lt 80", {"price": {"lt": 80}}, {}),
    ("F34", "6. Price filters", "price - lte 55", {"price": {"lte": 55}}, {}),
    ("F35", "6. Price filters", "price - range gte 50 lte 100", {"price": {"gte": 50, "lte": 100}}, {}),
    ("F36", "6. Price filters", "price - range gt 55 lt 100", {"price": {"gt": 55, "lt": 100}}, {}),
    ("F37", "7. MOQ filters", "moq - exact 2", {"moq": 2}, {}),
    ("F38", "7. MOQ filters", "moq - list [1, 3]", {"moq": [1, 3]}, {}),
    ("F39", "7. MOQ filters", "moq - gte 3", {"moq": {"gte": 3}}, {}),
    ("F40", "7. MOQ filters", "moq - range gt 1 lte 10", {"moq": {"gt": 1, "lte": 10}}, {}),
    ("F41", "8. Lead-time filters", "lead_time - exact 3", {"lead_time": 3}, {}),
    ("F42", "8. Lead-time filters", "lead_time - list [0, 8]", {"lead_time": [0, 8]}, {}),
    ("F43", "8. Lead-time filters", "lead_time - lte 1", {"lead_time": {"lte": 1}}, {}),
    ("F44", "8. Lead-time filters", "lead_time - gte 5", {"lead_time": {"gte": 5}}, {}),
    ("F45", "8. Lead-time filters", "lead_time - range gte 3 lt 5", {"lead_time": {"gte": 3, "lt": 5}}, {}),
    ("F46", "9. Supplier-count filters", "supplier_count - exact 0", {"supplier_count": 0}, {}),
    ("F47", "9. Supplier-count filters", "supplier_count - list [2, 3]", {"supplier_count": [2, 3]}, {}),
    ("F48", "9. Supplier-count filters", "supplier_count - gte 2", {"supplier_count": {"gte": 2}}, {}),
    ("F49", "9. Supplier-count filters", "supplier_count - lt 2", {"supplier_count": {"lt": 2}}, {}),
    # Boolean filters
    ("F50", "10. Boolean filters", "contract_indicator - false", {"contract_indicator": False}, {}),
    ("F51", "10. Boolean filters", "contract_indicator - true (no match)", {"contract_indicator": True}, {}),
    ("F52", "10. Boolean filters", "temporary_unavailable - true", {"temporary_unavailable": True}, {}),
    ("F53", "10. Boolean filters", "temporary_unavailable - false", {"temporary_unavailable": False}, {}),
    ("F54", "10. Boolean filters", "temporary_unavailable - list [true, false]", {"temporary_unavailable": [True, False]}, {}),
    # Combinations with sort / pagination / channel
    ("F55", "11. Combined filters", "category + price range + seller (AND)", {"category": [CAT_IT], "price": {"gte": 50, "lte": 60}, "seller": [SELLER_GCC]}, {}),
    ("F56", "11. Combined filters", "material_type + currency + temporary_unavailable", {"material_type": ["9CAT"], "currency": "SAR", "temporary_unavailable": False}, {}),
    ("F57", "11. Combined filters", "filters with no matching product", {"category": [CAT_ELEC], "material_type": "9COM"}, {}),
    ("F58", "11. Combined filters", "price range + sort by price asc", {"price": {"gte": 50, "lte": 100}}, {"sort_by": "price", "sort_order": "asc"}),
    ("F59", "11. Combined filters", "price range + sort by price desc", {"price": {"gte": 50, "lte": 100}}, {"sort_by": "price", "sort_order": "desc"}),
    ("F60", "11. Combined filters", "material_type + sort_by relevance", {"material_type": "9COM"}, {"sort_by": "relevance"}),
    ("F61", "11. Combined filters", "availability + pagination page 1 size 5", {"availability": "IN_STOCK"}, {"page": "1", "page_size": "5"}),
    ("F62", "11. Combined filters", "availability + pagination page 2 size 5", {"availability": "IN_STOCK"}, {"page": "2", "page_size": "5"}),
    ("F63", "11. Combined filters", "seller filter - mobile channel", {"seller": [SELLER_GCC]}, {"channel": "mobile"}),
    # Rejected filters
    ("E01", "12. Filter errors", "unknown filter key", {"brand": "HPE"}, {}),
    ("E02", "12. Filter errors", "malformed filters JSON", "{category: [abc]}", {}),
    ("E03", "12. Filter errors", "filters is a JSON array, not an object", [{"category": CAT_IT}], {}),
    ("E04", "12. Filter errors", "empty filters object (nothing else sent)", {}, {}),
    ("E05", "12. Filter errors", "keyword filter - empty string", {"material_type": ""}, {}),
    ("E06", "12. Filter errors", "keyword filter - empty list", {"category": []}, {}),
    ("E07", "12. Filter errors", "keyword filter - number instead of string", {"sku_id": 123}, {}),
    ("E08", "12. Filter errors", "keyword filter - range object", {"currency": {"gte": "SAR"}}, {}),
    ("E09", "12. Filter errors", "numeric filter - string value", {"price": "80"}, {}),
    ("E10", "12. Filter errors", "numeric filter - boolean value", {"moq": True}, {}),
    ("E11", "12. Filter errors", "numeric filter - unsupported range operator", {"price": {"eq": 80}}, {}),
    ("E12", "12. Filter errors", "numeric filter - empty range object", {"lead_time": {}}, {}),
    ("E13", "12. Filter errors", "numeric filter - string inside range", {"supplier_count": {"gte": "2"}}, {}),
    ("E14", "12. Filter errors", "numeric filter - list with a string", {"price": [80, "90"]}, {}),
    ("E15", "12. Filter errors", "boolean filter - string 'true'", {"contract_indicator": "true"}, {}),
    ("E16", "12. Filter errors", "boolean filter - number 1", {"temporary_unavailable": 1}, {}),
    ("E17", "12. Filter errors", "first invalid key is reported (valid + invalid)", {"category": [CAT_IT], "color": "red"}, {}),
]


def run_cases():
    out = []
    for cid, folder, name, filters, extra in CASES:
        form = {"channel": "web", "language": "en",
                "filters": filters if isinstance(filters, str) else json.dumps(filters),
                "page": "1", "page_size": "60", **extra}
        try:
            status, body = 200, search(cid, form)
        except ApiError as e:
            status, body = e.status, e.body
        out.append({"id": cid, "folder": folder, "name": f"{cid} {name}", "form": form, "status": status, "body": body})
    return out


# ---- writers

def write_xml(results, path):
    lines = ['<?xml version="1.0" encoding="UTF-8"?>',
             '<ApiSamples api="ASMO AI Search API" endpoint="POST /v1/search" contentType="multipart/form-data"'
             ' dataSource="api_samples/sample_products.json">',
             '  <Notes>Request fields are multipart/form-data parts; filters is a JSON string. '
             'Responses are computed over the products in sample_products.json. '
             'Every request also needs the header Authorization: Bearer &lt;Keycloak access token&gt;.</Notes>']
    folder = None
    for r in results:
        if r["folder"] != folder:
            if folder:
                lines.append("  </Group>")
            folder = r["folder"]
            lines.append(f'  <Group name="{escape(folder)}">')
        lines.append(f'    <Sample id="{r["id"]}" name="{escape(r["name"])}">')
        lines.append('      <Request method="POST" url="{{base_url}}/v1/search">')
        lines.append('        <Header name="Authorization">Bearer {{access_token}}</Header>')
        lines.append('        <Body contentType="multipart/form-data">')
        for k, v in r["form"].items():
            lines.append(f'          <Field name="{k}">{escape(v)}</Field>')
        lines.append("        </Body>")
        lines.append("      </Request>")
        lines.append(f'      <Response status="{r["status"]}" contentType="application/json">')
        body = json.dumps(r["body"], indent=2, ensure_ascii=False).replace("]]>", "]]]]><![CDATA[>")
        lines.append(f"        <Body><![CDATA[{body}]]></Body>")
        lines.append("      </Response>")
        lines.append("    </Sample>")
    lines += ["  </Group>", "</ApiSamples>", ""]
    path.write_text("\n".join(lines), encoding="utf-8")


def _tests(r):
    t = [f"test('Status {r['status']}', () => expect(res.getStatus()).to.equal({r['status']}));", "const b = res.getBody();"]
    if r["status"] != 200:
        t.append(f"test('code {r['body']['code']}', () => expect(b.code).to.eql('{r['body']['code']}'));")
        return "\n".join(t)
    f = json.loads(r["form"]["filters"])
    t.append("test('response shape', () => { ['request_id','total_results','page','page_size','products','facets'].forEach(k => expect(b).to.have.property(k)); });")
    t.append(f"test('page = {r['form']['page']}', () => expect(b.page).to.eql({r['form']['page']}));")
    if r["body"]["total_results"] == 0:
        t.append("test('no products', () => expect(b.products).to.have.lengthOf(0));")
    if "sku_id" in f:
        skus = f["sku_id"] if isinstance(f["sku_id"], list) else [f["sku_id"]]
        t.append(f"test('only requested SKUs', () => b.products.forEach(p => expect({json.dumps(skus)}).to.include(p.sku_id)));")
    for key in ("category", "category_name", "category_ids"):
        if key in f:
            ids = f[key] if isinstance(f[key], list) else [f[key]]
            t.append(f"test('category facet within filter', () => b.facets.category.forEach(c => expect({json.dumps(ids)}).to.include(c.id)));")
    for key in ("seller", "sellername"):
        if key in f:
            ids = f[key]
            t.append(f"test('every product has a requested seller', () => b.products.forEach(p => expect(p.seller_details.some(s => {json.dumps(ids)}.includes(s.seller_id))).to.be.true));")
    if r["form"].get("sort_by") == "price":
        cmp = "<=" if r["form"]["sort_order"] == "asc" else ">="
        t.append("test('sorted by primary seller price', () => { const p = b.products.map(x => x.seller_details[0] && x.seller_details[0].price).filter(x => x != null);"
                 f" for (let i = 1; i < p.length; i++) expect(p[i - 1] {cmp} p[i]).to.be.true; }});")
    return "\n".join(t)


def write_bruno(results, path):
    folders: dict[str, list] = {}
    for r in results:
        form = [{"type": "text", "name": k, "value": v, "enabled": True} for k, v in r["form"].items()]
        docs = ("Sample response (computed over api_samples/sample_products.json):\n\n"
                f"Status: {r['status']}\n\n```json\n{json.dumps(r['body'], indent=2, ensure_ascii=False)}\n```")
        items = folders.setdefault(r["folder"], [])
        items.append({
            "type": "http-request", "name": r["name"], "seq": len(items) + 1,
            "request": {
                "url": "{{base_url}}/v1/search", "method": "POST",
                "headers": [{"name": "Authorization", "value": "Bearer {{access_token}}", "enabled": True}],
                "params": [],
                "body": {"mode": "multipartForm", "json": None, "text": None, "xml": None, "sparql": None,
                         "formUrlEncoded": [], "multipartForm": form, "file": []},
                "script": {"req": "", "res": ""}, "vars": {"req": [], "res": []}, "assertions": [],
                "tests": _tests(r), "docs": docs, "auth": {"mode": "none"},
            },
        })
    env_vars = [("base_url", "http://localhost:8000", False), ("access_token", "", True)]
    collection = {
        "name": "ASMO AI Search API - Filter Samples",
        "version": "1",
        "items": [{"type": "folder", "name": name, "seq": i + 1, "items": items} for i, (name, items) in enumerate(folders.items())],
        "environments": [{"name": "local", "variables": [
            {"name": n, "value": v, "enabled": True, "secret": s, "type": "text"} for n, v, s in env_vars]}],
        "root": {
            "request": {"vars": {"req": [{"name": "base_url", "value": "http://localhost:8000", "enabled": True},
                                         {"name": "access_token", "value": "", "enabled": True}]}},
            "docs": "Every /v1/search filter key and value form (single value, list, range gt/gte/lt/lte, boolean), "
                    "combined with sort and pagination, plus every rejected filter. Each request's Docs tab holds the "
                    "response computed over api_samples/sample_products.json. Set access_token to a Keycloak access token.",
        },
        "brunoConfig": {"version": "1", "name": "ASMO AI Search API - Filter Samples", "type": "collection",
                        "ignore": ["node_modules", ".git"]},
    }
    path.write_text(json.dumps(collection, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")


def _compact_json(body):
    """JSON with one product / facet bucket per line, so long responses stay readable in the document."""
    def dump(v):
        return json.dumps(v, ensure_ascii=False)
    if "products" not in body:
        return json.dumps(body, indent=2, ensure_ascii=False)
    lines = ["{"]
    for k in ("request_id", "total_results", "page", "page_size"):
        lines.append(f'  "{k}": {dump(body[k])},')
    lines.append('  "products": [' + ("" if body["products"] else "],"))
    for i, p in enumerate(body["products"]):
        lines.append("    " + dump(p) + ("," if i < len(body["products"]) - 1 else ""))
    if body["products"]:
        lines.append("  ],")
    lines.append('  "facets": {')
    keys = list(body["facets"])
    for j, k in enumerate(keys):
        vals = body["facets"][k]
        end = "," if j < len(keys) - 1 else ""
        if not vals:
            lines.append(f'    "{k}": []{end}')
            continue
        lines.append(f'    "{k}": [')
        lines += ["      " + dump(v) + ("," if i < len(vals) - 1 else "") for i, v in enumerate(vals)]
        lines.append(f"    ]{end}")
    lines += ["  }", "}"]
    return "\n".join(lines)


FILTER_REFERENCE = [
    ("category / category_name / category_ids", "string or list of strings", "category id (facet id)", "category_ids"),
    ("sub_category / sub_category_name / sub_category_id", "string or list of strings", "sub-category id (facet id)", "sub_category_id"),
    ("product_type", "string or list of strings", "product-type id (facet id) or product-type name", "product_type_id or product_type"),
    ("product_type_id", "string or list of strings", "product-type id", "product_type_id"),
    ("seller / sellername", "string or list of strings", "seller id (facet id)", "any seller_details[].seller_id"),
    ("sold_by", "string or list of strings", "primary seller name", "sold_by"),
    ("material_type", "string or list of strings", "e.g. 9CAT, 9COM", "material_type"),
    ("availability", "string or list of strings", "e.g. IN_STOCK", "availability"),
    ("currency", "string or list of strings", "ISO 4217 code, e.g. SAR, USD", "primary seller currency"),
    ("sku_id / product_key / asmo_material_number / aramco_material_number", "string or list of strings", "exact value", "same field"),
    ("price", "number, list of numbers, or range {gt, gte, lt, lte}", "unit price", "primary seller price"),
    ("moq", "number, list of numbers, or range", "minimum order quantity", "primary seller moq"),
    ("lead_time", "number, list of numbers, or range", "days", "any seller_details[].lead_time"),
    ("supplier_count", "number, list of numbers, or range", "number of sellers", "supplier_count"),
    ("contract_indicator", "true / false (or list)", "boolean", "contract_indicator"),
    ("temporary_unavailable", "true / false (or list)", "boolean", "primary seller temporary_unavailable"),
]


def write_docx(results, path):
    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.shared import Pt

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)

    def table(header, rows, widths=None):
        t = doc.add_table(rows=1, cols=len(header))
        t.style = "Light Grid Accent 1"
        for c, h in zip(t.rows[0].cells, header):
            c.text = h
        for row in rows:
            for c, v in zip(t.add_row().cells, row):
                c.text = str(v)
        return t

    def code(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        for i, line in enumerate(text.split("\n")):
            run = p.add_run(line)
            run.font.name, run.font.size = "Consolas", Pt(7.5)
            if i < text.count("\n"):
                run.add_break(WD_BREAK.LINE)

    doc.add_heading("ASMO AI Search API — Filter Request / Response Samples", 0)
    doc.add_paragraph(
        "Sample requests and responses for every filter of POST /v1/search: each filter key, each accepted "
        "value form (single value, list, range with gt / gte / lt / lte, boolean), filters combined with sort, "
        "pagination and channel, and every rejected filter with its error. The same samples are in "
        "ASMO_Search_API_Filter_Samples.xml and in the Bruno collection "
        "ASMO_Search_API_Filter_Samples.bruno_collection.json.")

    doc.add_heading("1. Request", 1)
    doc.add_paragraph("POST {{base_url}}/v1/search — Content-Type: multipart/form-data — "
                      "Header: Authorization: Bearer <Keycloak access token>")
    table(["Field", "Required", "Values"], [
        ("channel", "yes", "web or mobile"),
        ("language", "yes", "en"),
        ("filters", "no", "JSON object sent as a string, e.g. {\"price\": {\"gte\": 50, \"lte\": 150}}"),
        ("sort_by", "no", "relevance (default) or price"),
        ("sort_order", "no", "desc (default) or asc"),
        ("page", "no", "1-based, default 1"),
        ("page_size", "no", "1–200, default 60"),
    ])
    doc.add_paragraph()
    doc.add_paragraph("At least one of search_query, image, filters or request_id is required. All filter keys "
                      "are combined with AND; a list value matches any of its values.")

    doc.add_heading("2. Filter keys", 1)
    table(["Filter key", "Value types", "Value", "Matches"], FILTER_REFERENCE)
    doc.add_paragraph()
    doc.add_paragraph("Any other key returns 400 INVALID_FILTER_KEY. A value of the wrong type (empty string, empty "
                      "list, string for a numeric key, number for a boolean key, range on a non-numeric key, unknown "
                      "range operator) returns 400 INVALID_FILTER_VALUE. filters that is not a JSON object returns "
                      "400 INVALID_FILTER.")

    doc.add_heading("3. Sample data", 1)
    doc.add_paragraph(f"The responses below are computed over these {len(PRODUCTS)} products from the "
                      "TejaAmbati-02/Test_Data repository (out_transformed folder), saved in "
                      "api_samples/sample_products.json. Between them they cover every value of every filter "
                      "field in that data.")
    table(["sku_id", "Product", "Category", "Material", "Price", "Currency", "MOQ", "Suppliers", "Temp. unavail.", "Sellers (lead time)"], [
        (d["sku_id"], d["product_name"], d["category_name"] or "—", d["material_type"],
         d["price"] if d["price"] is not None else "—", d["currency"] or "—",
         d["moq"] if d["moq"] is not None else "—", d["supplier_count"],
         d["temporary_unavailable"] if d["temporary_unavailable"] is not None else "—",
         ", ".join(f"{s['seller_name']} ({s['lead_time'] if s['lead_time'] is not None else '—'})"
                   for s in d["seller_details"] or []) or "—")
        for d in PRODUCTS])

    doc.add_heading("4. Summary of samples", 1)
    table(["ID", "Scenario", "filters", "Status", "Result"], [
        (r["id"], r["name"].split(" ", 1)[1], r["form"]["filters"], r["status"],
         f"{r['body']['total_results']} products" if r["status"] == 200 else r["body"]["code"])
        for r in results])

    doc.add_heading("5. Samples", 1)
    folder = None
    for r in results:
        if r["folder"] != folder:
            folder = r["folder"]
            doc.add_heading(folder, 2)
        doc.add_heading(r["name"], 3)
        doc.add_paragraph("Request — POST {{base_url}}/v1/search (multipart/form-data)")
        table(["Form field", "Value"], list(r["form"].items()))
        doc.add_paragraph()
        doc.add_paragraph(f"Response — HTTP {r['status']}")
        code(_compact_json(r["body"]))
    doc.save(path)


if __name__ == "__main__":
    results = run_cases()
    write_xml(results, HERE / "ASMO_Search_API_Filter_Samples.xml")
    write_bruno(results, HERE / "ASMO_Search_API_Filter_Samples.bruno_collection.json")
    write_docx(results, HERE / "ASMO_Search_API_Filter_Samples.docx")
    for r in results:
        print(f"{r['name']:<70} {r['status']} {r['body'].get('total_results', r['body'].get('code'))}")
